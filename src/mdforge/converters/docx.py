"""DOCX converter (powered by python-docx).

Renders paragraphs, headings, nested lists, tables and inline formatting in
document order by walking the underlying OOXML body - which is what makes the
output usable rather than a pile of text boxes.
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from typing import Any

from ..core.base import BaseConverter, DocumentResult, PathLike
from ..exceptions import ConversionError
from ..utils.deps import import_optional
from ..utils.text import markdown_table

__all__ = ["DocxConverter"]

_HEADING_RE = re.compile(r"^heading\s*(\d+)$", re.IGNORECASE)
_CODE_STYLES = {"code", "macro text", "html code", "source code", "html preformatted"}


class DocxConverter(BaseConverter):
    """Convert Word documents into structured Markdown.

    Supported options:

    ``max_table_rows``
        Truncate tables (default ``100``) to bound prompt size.
    ``include_headers_footers``
        Include section header/footer text (default ``False``).
    """

    name = "docx"
    extensions = (".docx",)
    mime_types = ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",)
    requires = ("docx",)
    extra = "docx"
    description = "Word documents (.docx)"

    def convert(self, path: PathLike, **options: Any) -> DocumentResult:
        """Convert ``path`` to Markdown.

        Args:
            path: ``.docx`` file to convert.
            **options: See the class docstring.

        Returns:
            A :class:`DocumentResult` with the document body.

        Raises:
            MissingDependencyError: If ``python-docx`` is not installed.
            ConversionError: If the file is not a readable Word document.
        """
        docx = import_optional("docx", extra="docx", feature="to convert Word documents")

        try:
            document = docx.Document(str(path))
        except Exception as exc:
            raise ConversionError(
                str(path),
                f"cannot open Word document (is it a legacy .doc?): {exc}",
                converter=self.name,
            ) from exc

        max_table_rows = int(options.get("max_table_rows", 100))
        warnings: list[str] = []
        heading_count = 0
        table_count = 0
        image_count = len(getattr(document, "inline_shapes", []) or [])

        items: list[tuple[str, str, int]] = []

        for block in _iter_block_items(document):
            if block.__class__.__name__ == "Paragraph":
                rendered = _render_paragraph(block)
                if rendered is None:
                    continue
                if rendered[1] == "heading":
                    heading_count += 1
                items.append(rendered)
            else:
                rows = _table_rows(block, max_rows=max_table_rows)
                if rows:
                    table_count += 1
                    if len(rows) > max_table_rows:
                        warnings.append(f"A table was truncated to {max_table_rows} rows.")
                    items.append((markdown_table(rows[:max_table_rows]), "block", 0))

        body = _assemble(items)

        if options.get("include_headers_footers"):
            for section in document.sections:
                for part_name, part in (
                    ("Header", section.header),
                    ("Footer", section.footer),
                ):
                    text = "\n".join(p.text for p in part.paragraphs if p.text.strip())
                    if text.strip():
                        body = f"{body}\n\n<!-- {part_name} -->\n\n{text.strip()}"

        if image_count:
            warnings.append(
                f"Document contains {image_count} inline image(s); "
                "images are not transcribed (use an OCR/vision pipeline for that)."
            )

        metadata = _core_properties(document)
        metadata.update(
            {
                "table_count": table_count,
                "image_count": image_count,
                "heading_count": heading_count,
                "paragraph_count": len(document.paragraphs),
            }
        )
        return self.build_result(path, body, metadata=metadata, warnings=warnings)


def _iter_block_items(parent: Any) -> Iterator[Any]:
    """Yield paragraphs and tables from ``parent`` in document order.

    ``python-docx`` exposes paragraphs and tables as separate collections,
    which loses ordering.  Walking the OOXML body restores it.
    """
    from docx.document import Document as _Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph

    if isinstance(parent, _Document):
        parent_element = parent.element.body
    elif isinstance(parent, _Cell):
        parent_element = parent._tc
    else:  # pragma: no cover - defensive
        raise TypeError(f"Unsupported parent type: {type(parent)!r}")

    for child in parent_element.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _render_paragraph(paragraph: Any) -> tuple[str, str, int] | None:
    """Render one paragraph to ``(markdown, kind, indent)``.

    ``kind`` is ``"heading"``, ``"list"`` or ``"block"``; ``indent`` is the
    list nesting depth and is only non-zero for list items.  Returns ``None``
    for paragraphs that produce no output at all.
    """
    text = _paragraph_inline(paragraph).strip()
    style_name = (getattr(paragraph.style, "name", "") or "").strip()
    lowered = style_name.lower()

    match = _HEADING_RE.match(style_name)
    if match:
        level = max(1, min(6, int(match.group(1))))
        return (f"{'#' * level} {text}", "heading", 0) if text else None
    if lowered == "title":
        return (f"# {text}", "heading", 0) if text else None
    if lowered == "subtitle":
        return (f"**{text}**", "block", 0) if text else None
    if lowered.startswith("list") or lowered == "list paragraph":
        if not text:
            return None
        marker = "1." if "number" in lowered else "-"
        return (f"{marker} {text}", "list", _list_level(style_name, paragraph))
    if "quote" in lowered:
        return ("\n".join(f"> {line}" for line in text.split("\n")), "block", 0) if text else None
    if lowered in _CODE_STYLES:
        return (f"`{text}`", "block", 0) if text else None
    return (text, "block", 0) if text else None


def _assemble(items: list[tuple[str, str, int]]) -> str:
    """Join rendered blocks into a document, keeping lists tight.

    Consecutive list items are joined with a single newline (a "tight" list);
    everything else is separated by a blank line.  Without this, every bullet
    in a Word document becomes its own loose list.
    """
    assembled: list[tuple[str, str]] = []
    for text, kind, indent in items:
        if not text.strip():
            continue
        markdown = _indent(text, indent) if kind == "list" else text
        if kind == "list" and assembled and assembled[-1][1] == "list":
            assembled[-1] = (f"{assembled[-1][0]}\n{markdown}", "list")
        else:
            assembled.append((markdown, kind))
    return "\n\n".join(markdown for markdown, _ in assembled)


def _list_level(style_name: str, paragraph: Any) -> int:
    """Best-effort nesting depth for a list paragraph."""
    digits = re.findall(r"(\d+)", style_name)
    if digits:
        return max(0, int(digits[-1]) - 1)
    indent = getattr(paragraph.paragraph_format, "left_indent", None)
    if indent is not None:
        try:
            steps = int(indent.pt // 18)
            return max(0, min(5, steps))
        except (AttributeError, TypeError):  # pragma: no cover - defensive
            return 0
    return 0


def _paragraph_inline(paragraph: Any) -> str:
    """Render a paragraph's runs and hyperlinks in document order."""
    parts: list[str] = []
    for child in paragraph._p.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "r":
            parts.append(_render_run_element(child))
        elif tag == "hyperlink":
            parts.append(_render_hyperlink(child, paragraph))
    return "".join(parts)


def _render_hyperlink(element: Any, paragraph: Any) -> str:
    """Render a ``w:hyperlink`` element, resolving its relationship target."""
    text = "".join(
        _render_run_element(run) for run in element.iterchildren() if _local(run.tag) == "r"
    )
    if not text:
        return ""
    relationship_id = None
    for key, value in element.attrib.items():
        if _local(key) == "id":
            relationship_id = value
            break
    if not relationship_id:
        return text
    try:
        address = paragraph.part.rels[relationship_id].target_ref
    except (KeyError, AttributeError):  # pragma: no cover - malformed doc
        return text
    if str(address).startswith("#"):
        return text
    return f"[{text}]({address})"


def _render_run_element(element: Any) -> str:
    """Render a ``w:r`` run element including its character formatting."""
    text = _run_text(element)
    if not text:
        return ""
    spec = _run_format(element)
    if spec["code"]:
        text = f"`{text}`"
    if spec["bold"] and spec["italic"]:
        return f"***{text}***"
    if spec["bold"]:
        return f"**{text}**"
    if spec["italic"]:
        return f"*{text}*"
    if spec["strike"]:
        return f"~~{text}~~"
    return text


def _run_text(element: Any) -> str:
    """Concatenate the text/breaks/tabs of a run element."""
    chunks: list[str] = []
    for node in element.iter():
        tag = _local(node.tag)
        if tag == "t" and node.text:
            chunks.append(node.text)
        elif tag == "tab":
            chunks.append("\t")
        elif tag in {"br", "cr"}:
            chunks.append("\n")
    return "".join(chunks)


def _run_format(element: Any) -> dict[str, bool]:
    """Read bold/italic/strike/code flags from a run's ``rPr``."""
    spec = {"bold": False, "italic": False, "strike": False, "code": False}
    properties = None
    for child in element.iterchildren():
        if _local(child.tag) == "rPr":
            properties = child
            break
    if properties is None:
        return spec
    for child in properties.iterchildren():
        tag = _local(child.tag)
        value = child.get(_qn_attr(child, "val"))
        enabled = value not in {"0", "false", "off", "none"}
        if tag == "b":
            spec["bold"] = enabled
        elif tag == "i":
            spec["italic"] = enabled
        elif tag == "strike":
            spec["strike"] = enabled
        elif tag in {"rStyle"} and (value or "").lower() in _CODE_STYLES:
            spec["code"] = True
    style = properties.find(_qn("w:rStyle"))
    if style is not None and (style.get(_qn("w:val")) or "").lower() in _CODE_STYLES:
        spec["code"] = True
    return spec


def _table_rows(table: Any, *, max_rows: int) -> list[list[str]]:
    """Extract a table as a list of rows of cell text."""
    rows: list[list[str]] = []
    for index, row in enumerate(table.rows):
        if index > max_rows:
            break
        cells: list[str] = []
        for cell in row.cells:
            paragraphs = [_paragraph_inline(p).strip() for p in cell.paragraphs if p.text.strip()]
            cells.append("<br>".join(p for p in paragraphs if p))
        if any(cells):
            rows.append(cells)
    return rows


def _core_properties(document: Any) -> dict[str, Any]:
    """Extract and normalise ``docx`` core properties."""
    props = getattr(document, "core_properties", None)
    if props is None:  # pragma: no cover - defensive
        return {}
    mapping = {
        "title": "title",
        "author": "author",
        "subject": "subject",
        "category": "category",
        "comments": "description",
        "keywords": "keywords",
        "last_modified_by": "last_modified_by",
        "revision": "revision",
    }
    metadata: dict[str, Any] = {}
    for attribute, key in mapping.items():
        value = getattr(props, attribute, None)
        if value is None:
            continue
        if key == "revision":
            metadata[key] = int(value)
            continue
        text = str(value).strip()
        if text:
            metadata[key] = text
    for attribute, key in (("created", "created"), ("modified", "modified")):
        value = getattr(props, attribute, None)
        if value is not None:
            metadata[key] = value.isoformat()
    return metadata


def _local(tag: str) -> str:
    """Return the local part of a namespaced OOXML tag."""
    return tag.rsplit("}", 1)[-1]


def _qn(tag: str) -> str:
    """Qualify a ``w:`` prefixed tag using python-docx's namespace map."""
    from docx.oxml.ns import qn

    return str(qn(tag))


def _qn_attr(element: Any, name: str) -> str:
    """Return the fully-qualified attribute name for ``name`` on ``element``."""
    for key in element.attrib:
        if _local(key) == name:
            return str(key)
    return _qn(f"w:{name}")


def _indent(text: str, level: int) -> str:
    """Indent a rendered list item according to its nesting level."""
    if level <= 0 or not text:
        return text
    prefix = "  " * level
    return "\n".join(f"{prefix}{line}" if line.strip() else line for line in text.split("\n"))
